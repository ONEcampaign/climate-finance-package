"""The BTR reader on stripped real Table III.1 layouts, the ingest and party weights."""

import json
import shutil
from pathlib import Path

import openpyxl
import polars as pl
import pytest
import yaml
from typer.testing import CliRunner

from climate_finance import schema
from climate_finance.build.cli import app
from climate_finance.build.normalise import btr_at_dac_rates
from climate_finance.build.readers.btr import (
    estimate_party_weights,
    ingest_unfccc,
    party_provider,
    read_btr_archive,
    read_btr_workbook,
    write_party_weights,
)
from climate_finance.errors import SourceUnavailableError

BTR = Path(__file__).parent / "fixtures" / "raw" / "btr"

# Per fixture, each kept row: amount_domestic, face_usd, grant_equiv_usd, status.
LAYOUTS = {
    "usa-millions-face-value-row-under-currency": (
        "USA-table-1.xlsx",
        {
            "BTR-USA-2021-00009": (None, 0.5e6, None, "committed"),
            "BTR-USA-2021-00010": (None, 2e6, 1.5e6, "committed"),
            # USD left empty, filled at the stated rate of 1.
            "BTR-USA-2022-00009": (0.4e6, 0.4e6, None, "committed"),
            "BTR-USA-2022-00010": (1.2e6, 1.2e6, None, "committed"),
        },
    ),
    "gbr-millions-rate-as-dollar-to-pound": (
        "GBR-table-1.xlsx",
        {
            "BTR-GBR-2021-00010": (1.454e6, 2e6, 2e6, "committed"),
            "BTR-GBR-2021-00011": (0.727e6, 1e6, 1e6, "committed"),
        },
    ),
    "can-xlsm-in-zip-with-year-column": (
        "CAN-tables.zip",
        {
            "BTR-CAN-2021-00006": (1e6, 0.8e6, None, "disbursed"),
            "BTR-CAN-2022-00007": (0.5e6, 0.4e6, None, "disbursed"),
        },
    ),
    "bel-rate-in-label-cell-and-provided-status": (
        "BEL-table-1.xlsx",
        {
            "BTR-BEL-2021-00010": (881_000, 1e6, 1e6, "disbursed"),
            "BTR-BEL-2021-00011": (88_100, 1e5, 1e5, "disbursed"),
        },
    ),
    "deu-other-year-and-total-rows-dropped": (
        "DEU-table-1.xlsx",
        {
            "BTR-DEU-2021-00010": (850_000, 1e6, None, "committed"),
            "BTR-DEU-2021-00011": (1.7e6, 2e6, None, "committed"),
        },
    ),
    "dnk-committed-and-disbursed-both-kept": (
        "DNK-table-1.xlsx",
        {
            "BTR-DNK-2021-00010": (633_870, 1e5, None, "committed"),
            "BTR-DNK-2021-00011": (316_935, 5e4, None, "disbursed"),
        },
    ),
    "fra-usd-filled-at-sheet-rate": (
        "FRA-table-1.xlsx",
        {
            "BTR-FRA-2021-00010": (845_500, 1e6, 5e5, "committed"),
            "BTR-FRA-2021-00011": (1.691e6, 2e6, None, "committed"),
        },
    ),
    "aus-grant-equivalent-only": (
        "AUS-table-1.xlsx",
        {
            "BTR-AUS-2021-00010": (None, None, 1e6, "disbursed"),
            "BTR-AUS-2021-00011": (None, None, 5e5, "disbursed"),
        },
    ),
    "fin-title-only-and-ua-rows-dropped": (
        "FIN-table-1.xlsx",
        {
            "BTR-FIN-2021-00010": (845_600, 1e6, 1e6, "disbursed"),
            "BTR-FIN-2021-00012": (422_800, 5e5, 5e5, "disbursed"),
            "BTR-FIN-2021-00013": (169_120, 2e5, 2e5, "disbursed"),
        },
    ),
    "mco-footnote-number-on-amount": (
        "MCO-table-1.xlsx",
        {"BTR-MCO-2021-00010": (100_000, 117_647.0588, 117_647.0588, "disbursed")},
    ),
    "chl-decimal-comma-and-spanish-status": (
        "CHL-table-1.xlsx",
        {"BTR-CHL-2023-00010": (990_000.5, 1_000.25, None, None)},
    ),
    "idn-year-in-additional-information": (
        "IDN-table-1.xlsx",
        {
            "BTR-IDN-2021-00007": (None, None, 1e5, "disbursed"),
            "BTR-IDN-2022-00008": (None, None, 2e5, "disbursed"),
        },
    ),
    "hun-rate-with-decimal-comma": (
        "HUN-table-1.xlsx",
        {"BTR-HUN-2021-00010": (30_329_000, 1e5, 1e5, "disbursed")},
    ),
    "nld-rate-as-dollars-to-the-euro": (
        "NLD-table-1.xlsx",
        {
            "BTR-NLD-2021-00010": (
                1.183e6,
                1.183e6 * 1.183,
                1.183e6 * 1.183,
                "disbursed",
            )
        },
    ),
}


def _read(name: str) -> pl.DataFrame:
    return read_btr_workbook(BTR / name, party=name[:3])


@pytest.mark.parametrize(("name", "expected"), LAYOUTS.values(), ids=LAYOUTS.keys())
def test_reads_each_layout(name: str, expected: dict[str, tuple]) -> None:
    frame = _read(name)

    schema.validate(schema.BTR, frame)
    got = {
        row[0]: row[1:]
        for row in frame.select(
            "row_id", "amount_domestic", "face_usd", "grant_equiv_usd", "status"
        ).iter_rows()
    }
    assert got.keys() == expected.keys()
    for row_id, values in expected.items():
        assert got[row_id] == pytest.approx(values), row_id


def test_finland_aggregates_are_flagged_and_na_lines_leave_the_recipient() -> None:
    frame = _read("FIN-table-1.xlsx")

    flagged = frame.filter(pl.col("is_aggregate"))["title"].sort().to_list()
    assert flagged == [
        "Finnfund climate-related investments",
        "Other bilateral projects",
    ]
    assert frame["recipient"].to_list() == ["Africa", "Global", "Global"]


def test_rows_keep_title_party_currency_and_recipient_code() -> None:
    frame = _read("GBR-table-1.xlsx")

    assert frame["title"].to_list() == ["Activity A", "201575 Programme B"]
    assert set(frame["currency"]) == {"GBP"}
    # Exact recipient names get their DAC code; "Global" has none.
    codes = _read("FRA-table-1.xlsx")["recipient_code"].to_list()
    assert codes == [71, 610]
    assert frame["recipient_code"].is_null().all()


@pytest.mark.parametrize(
    ("party", "code"),
    [("DEU", 5), ("EUU", 918), ("ARM", None), ("CHN", None), ("KEN", None)],
)
def test_a_party_that_is_no_dac_provider_has_no_provider_code(
    party: str, code: int | None
) -> None:
    assert party_provider(party) == code


def test_rows_of_a_party_without_a_provider_code_are_kept() -> None:
    frame = read_btr_workbook(BTR / "FRA-table-1.xlsx", party="ARM")

    schema.validate(schema.BTR, frame)
    assert frame.height == 2
    assert frame["provider_code"].is_null().all()


@pytest.mark.parametrize(
    ("name", "party", "unit"),
    [("USA-table-1.xlsx", "DEU", "millions"), ("FRA-table-1.xlsx", "USA", "units")],
)
def test_amounts_contradicting_the_party_unit_are_refused(
    name: str, party: str, unit: str
) -> None:
    with pytest.raises(ValueError, match=f"read as {unit}"):
        read_btr_workbook(BTR / name, party=party)


def test_amounts_are_restated_at_the_dac_rate() -> None:
    # The United Kingdom converts at 0.727 pounds a dollar, against a DAC rate of
    # 0.8112; the United States leaves its 2021 domestic amounts empty.
    btr = pl.concat([_read("GBR-table-1.xlsx"), _read("USA-table-1.xlsx")])
    exchange = pl.DataFrame(
        {"iso_code": ["GBR", "USA"], "year": [2021, 2021], "exchange_rate": [0.8112, 1]}
    )

    restated = btr_at_dac_rates(btr, exchange=exchange)

    schema.validate(schema.BTR, restated)
    got = {
        row[0]: row[1:]
        for row in restated.select(
            "row_id", "dac_rate", "face_usd_at_dac", "grant_equiv_usd_at_dac"
        ).iter_rows()
    }
    assert got["BTR-GBR-2021-00011"] == pytest.approx(
        (0.8112, 0.727e6 / 0.8112, 0.727e6 / 0.8112)
    )
    assert got["BTR-USA-2021-00010"] == pytest.approx((1.0, 2e6, 1.5e6))
    # No DAC rate for the year: the party's own USD stands.
    assert got["BTR-USA-2022-00009"] == (None, 0.4e6, None)


def _edited(tmp_path: Path, name: str, sheet: str, cells: dict[str, object]) -> Path:
    workbook = openpyxl.load_workbook(BTR / name)
    for cell, value in cells.items():
        workbook[sheet][cell] = value
    path = tmp_path / name
    workbook.save(path)
    return path


@pytest.mark.parametrize(
    ("cells", "message"),
    [
        # France row 10 reports 845,500 EUR as 2,000,000 USD, implying 0.42 a dollar.
        ({"E10": 2_000_000}, "differs from the rate its rows imply"),
        ({"C3": None}, "no rate stated"),
        ({"H10": "Planned"}, "outside the vocabulary"),
    ],
    ids=["rate-contradicted", "rate-missing", "unknown-status"],
)
def test_unreadable_sheets_are_refused(
    tmp_path: Path, cells: dict[str, object], message: str
) -> None:
    path = _edited(tmp_path, "FRA-table-1.xlsx", "Table1_2021", cells)

    with pytest.raises(ValueError, match=message):
        read_btr_workbook(path, party="FRA")


def _archive(tmp_path: Path, files: dict[str, str]) -> Path:
    """An archive of fixtures under manifest names, with the manifest."""
    archive = tmp_path / "archive"
    archive.mkdir()
    entries = []
    for name, party in files.items():
        shutil.copy(BTR / name.rsplit("__", 1)[1], archive / name)
        size = (archive / name).stat().st_size
        entries.append({"party": party, "name": name, "ok": True, "bytes": size})
    (archive / "manifest.json").write_text(json.dumps(entries), encoding="utf-8")
    return archive


def test_ingest_writes_every_listed_workbook(tmp_path: Path) -> None:
    archive = _archive(
        tmp_path,
        {"Canada__CAN-tables.zip": "Canada", "France__FRA-table-1.xlsx": "France"},
    )
    out = tmp_path / "btr.parquet"

    result = CliRunner().invoke(app, ["ingest-unfccc", str(archive), "--out", str(out)])

    assert result.exit_code == 0, result.output
    btr = pl.read_parquet(out)
    schema.validate(schema.BTR, btr)
    assert btr["row_id"].to_list() == sorted(btr["row_id"])
    assert btr.group_by("party").len().sort("party").rows() == [("CAN", 2), ("FRA", 2)]


def test_ingest_records_each_workbook_hash(tmp_path: Path) -> None:
    archive = _archive(
        tmp_path,
        {"Canada__CAN-tables.zip": "Canada", "France__FRA-table-1.xlsx": "France"},
    )
    out = tmp_path / "btr.parquet"

    def hashes() -> dict[str, str]:
        ingest_unfccc(archive, out_path=out)
        manifest = pl.read_parquet(tmp_path / "btr.manifest.parquet")
        schema.validate(schema.RAW_MANIFEST, manifest)
        return dict(manifest.select("path", "sha256").iter_rows())

    before = hashes()
    france = archive / "France__FRA-table-1.xlsx"
    workbook = openpyxl.load_workbook(france)
    workbook["Table1_2021"]["B10"] = "ALGERIA"
    workbook.save(france)
    entries = json.loads((archive / "manifest.json").read_text(encoding="utf-8"))
    entries[1]["bytes"] = france.stat().st_size
    (archive / "manifest.json").write_text(json.dumps(entries), encoding="utf-8")
    after = hashes()

    assert (
        before.keys()
        == after.keys()
        == {
            str(archive / "Canada__CAN-tables.zip"),
            str(france),
        }
    )
    assert after[str(france)] != before[str(france)]
    canada = str(archive / "Canada__CAN-tables.zip")
    assert after[canada] == before[canada]


def test_ingest_refuses_a_file_unlike_its_manifest_entry(tmp_path: Path) -> None:
    archive = _archive(tmp_path, {"France__FRA-table-1.xlsx": "France"})
    manifest = archive / "manifest.json"
    entries = json.loads(manifest.read_text(encoding="utf-8"))
    entries[0]["bytes"] += 1
    manifest.write_text(json.dumps(entries), encoding="utf-8")

    with pytest.raises(SourceUnavailableError, match="manifest says"):
        ingest_unfccc(archive, out_path=tmp_path / "btr.parquet")


def test_ingest_refuses_two_sheets_for_one_party_year(tmp_path: Path) -> None:
    archive = _archive(
        tmp_path,
        {
            "France__FRA-table-1.xlsx": "France",
            "France__draft__FRA-table-1.xlsx": "France",
        },
    )

    with pytest.raises(ValueError, match="Several sheets give one party's year"):
        ingest_unfccc(archive, out_path=tmp_path / "btr.parquet")


def test_the_archive_counts_each_workbooks_rows_kept_and_dropped(
    tmp_path: Path,
) -> None:
    archive = _archive(
        tmp_path,
        {
            "Finland__FIN-table-1.xlsx": "Finland",
            "Germany__V0.1_started__DEU-table-1.xlsx": "Germany",
        },
    )
    # A party whose workbook holds support-received tables only.
    empty = openpyxl.Workbook()
    empty.active.title = "Table 7"
    empty.save(archive / "Kenya__KEN.xlsx")
    entries = json.loads((archive / "manifest.json").read_text(encoding="utf-8"))
    entries.append(
        {
            "party": "Kenya",
            "name": "Kenya__KEN.xlsx",
            "ok": True,
            "bytes": (archive / "Kenya__KEN.xlsx").stat().st_size,
        }
    )
    (archive / "manifest.json").write_text(json.dumps(entries), encoding="utf-8")

    read = read_btr_archive(archive)

    schema.validate(schema.BTR_WORKBOOKS, read.workbooks)
    counts = {row.pop("file"): row for row in read.workbooks.to_dicts()}
    germany = "Germany__V0.1_started__DEU-table-1.xlsx"
    assert counts == {
        # A title with no amount, and a row of "UA".
        "Finland__FIN-table-1.xlsx": {
            "party": "FIN",
            "sheets": ["Finland__FIN-table-1.xlsx#Table1_2021"],
            "rows_kept": 3,
            "dropped_blank": 1,
            "dropped_notation": 1,
            "dropped_total": 0,
            "draft_marker": None,
        },
        # The total row, and on 2022 an empty activity and one of "NA".
        germany: {
            "party": "DEU",
            "sheets": [f"{germany}#Table1_2021", f"{germany}#Table1_2022"],
            "rows_kept": 2,
            "dropped_blank": 1,
            "dropped_notation": 1,
            "dropped_total": 1,
            "draft_marker": "started",
        },
        "Kenya__KEN.xlsx": {
            "party": "KEN",
            "sheets": [],
            "rows_kept": 0,
            "dropped_blank": 0,
            "dropped_notation": 0,
            "dropped_total": 0,
            "draft_marker": None,
        },
    }
    assert read.rows.group_by("party").len().sort("party").rows() == [
        ("DEU", 2),
        ("FIN", 3),
    ]


def _pairs(party: str, rows: list[tuple[float, float, int]]) -> pl.DataFrame:
    return pl.DataFrame(
        [(party, *row) for row in rows],
        schema={
            "party": pl.String,
            "btr_usd": pl.Float64,
            "crs_usd": pl.Float64,
            "rio_marker": pl.Int8,
        },
        orient="row",
    )


def test_party_weight_is_the_level_of_most_significant_pairs() -> None:
    pairs = pl.concat(
        [
            # Norway's shape: Significant at 0.4, Principal at 1.0.
            _pairs("NOR", [(40, 100, 1), (81, 200, 1), (100, 100, 2), (101, 100, 2)]),
            # One Significant pair at each of 0.5 and 1.0: the lower level wins.
            _pairs("DEU", [(50, 100, 1), (100, 100, 1), (70, 100, 1), (5, 0, 1)]),
            # No pair at 1.0, so no share of them is Principal; no Significant
            # pair, so no share of those at a level.
            _pairs("AUT", [(50, 100, 1)]),
            _pairs("EST", [(100, 100, 2)]),
        ]
    )

    weights = {row["party"]: row for row in estimate_party_weights(pairs).to_dicts()}

    assert weights["NOR"]["significant"] == 0.4
    assert weights["NOR"]["share_at"] == {"0.4": 0.5, "0.5": 0.0, "1.0": 0.5}
    assert weights["NOR"]["principal_at_100"] == 1.0
    assert weights["DEU"]["significant"] == 0.5
    # A pair with no CRS amount has no ratio and is left out.
    assert weights["DEU"]["pairs"] == 3
    assert weights["DEU"]["significant_share_at"]["1.0"] == pytest.approx(1 / 3)
    assert weights["AUT"]["principal_at_100"] is None
    assert weights["EST"]["significant_share_at"]["1.0"] is None


def test_party_weights_are_written_unconfirmed(tmp_path: Path) -> None:
    path = tmp_path / "unfccc_party_weights.yaml"
    pairs = _pairs("NOR", [(40, 100, 1), (100, 100, 2)])

    write_party_weights(pairs, path=path)

    entry = yaml.safe_load(path.read_text(encoding="utf-8"))["NOR"]
    assert entry["significant"] == 0.4
    assert entry["confirmed"] is False
    assert entry["evidence"]["pairs"] == 2


def test_party_weights_keep_a_reviewer_confirmation(tmp_path: Path) -> None:
    path = tmp_path / "unfccc_party_weights.yaml"
    pairs = _pairs("NOR", [(40, 100, 1)])
    write_party_weights(pairs, path=path)
    path.write_text(path.read_text().replace("confirmed: false", "confirmed: true"))

    with pytest.raises(FileExistsError, match="NOR"):
        write_party_weights(pairs, path=path)
